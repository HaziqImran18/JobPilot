"""Resume upload, text extraction and profile import.

Resumes are stored inside the signed-in account's private uploads directory.
The extracted text is kept so the agent can reason about the CV without reading
the binary again. Extraction never invents facts: if a field cannot be found it
stays blank and the user can fill it in on the review step.
"""

from __future__ import annotations

import hashlib
import io
import re
import uuid
from contextlib import suppress
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from jobpilot.core.config import get_settings, user_settings
from jobpilot.core.db import Database, get_db
from jobpilot.core.user_context import current_user_id
from jobpilot.core.utils import json_dumps, json_loads, normalize_text, utc_now_iso
from jobpilot.pipelines.extract import extract_skills
from jobpilot.pipelines.locations import resolve_location

MAX_RESUME_BYTES = 8 * 1024 * 1024
ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}
_MAGIC = {
    ".pdf": (b"%PDF",),
    ".docx": (b"PK\x03\x04",),
    ".txt": (),
}
_EMAIL = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
_PHONE = re.compile(r"(?:\+?\d[\d\s().-]{7,}\d)")
_LINKEDIN = re.compile(r"(?:https?://)?(?:www\.)?linkedin\.com/in/[A-Za-z0-9_%.-]+", re.I)
_GITHUB = re.compile(r"(?:https?://)?(?:www\.)?github\.com/[A-Za-z0-9_.-]+", re.I)
_URL = re.compile(r"https?://[^\s)>\]]+", re.I)
_YEARS = re.compile(r"(\d{1,2})\+?\s*years?", re.I)
_HEADLINE_HINTS = (
    "engineer",
    "developer",
    "analyst",
    "manager",
    "designer",
    "consultant",
    "specialist",
    "scientist",
    "architect",
    "accountant",
    "administrator",
    "intern",
    "associate",
    "director",
    "lead",
    "officer",
    "executive",
    "teacher",
    "nurse",
    "marketer",
)


class ResumeError(Exception):
    """A resume operation failed in a way the UI should explain."""


@dataclass
class Resume:
    id: str
    filename: str
    mime_type: str = ""
    size_bytes: int = 0
    status: str = "uploaded"
    text: str = ""
    extracted: dict[str, Any] = field(default_factory=dict)
    error: str = ""
    is_active: bool = False
    created_at: str = ""
    updated_at: str = ""

    def summary(self) -> dict[str, Any]:
        """Metadata without the full extracted text."""
        return {
            "id": self.id,
            "filename": self.filename,
            "mime_type": self.mime_type,
            "size_bytes": self.size_bytes,
            "status": self.status,
            "is_active": self.is_active,
            "error": self.error,
            "extracted": self.extracted,
            "characters": len(self.text),
            "created_at": self.created_at,
            "updated_at": self.updated_at,
        }


def _new_id() -> str:
    return f"resume_{uuid.uuid4().hex[:12]}"


def _guess_extension(filename: str) -> str:
    return Path(filename or "").suffix.lower()


def _clean_filename(filename: str) -> str:
    name = Path(filename or "resume").name
    name = re.sub(r"[^A-Za-z0-9._ -]", "_", name).strip(" .")
    return name[:120] or "resume"


def _extract_pdf(data: bytes) -> tuple[str, str]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - dependency is declared
        raise ResumeError("PDF support is not installed on this server.") from exc
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as exc:  # noqa: BLE001 - reported to the user
                raise ResumeError("This PDF is password protected.") from exc
        pages = []
        for page in reader.pages:
            try:
                pages.append(page.extract_text() or "")
            except Exception:  # noqa: BLE001 - one broken page should not stop the rest
                continue
        return "\n".join(pages), ""
    except ResumeError:
        raise
    except Exception as exc:  # noqa: BLE001 - malformed PDFs vary widely
        raise ResumeError(f"Could not read this PDF: {exc}") from exc


def _extract_docx(data: bytes) -> tuple[str, str]:
    try:
        import docx
    except ImportError as exc:  # pragma: no cover - dependency is declared
        raise ResumeError("DOCX support is not installed on this server.") from exc
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001 - corrupted archives raise many types
        raise ResumeError(f"Could not read this DOCX: {exc}") from exc
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(dict.fromkeys(cells)))
    return "\n".join(parts), ""


def _extract_text(extension: str, data: bytes) -> tuple[str, str]:
    if extension == ".pdf":
        return _extract_pdf(data)
    if extension == ".docx":
        return _extract_docx(data)
    try:
        return data.decode("utf-8", errors="replace"), ""
    except Exception as exc:  # noqa: BLE001 - defensive
        raise ResumeError(f"Could not read this file: {exc}") from exc


def extract_fields(text: str) -> dict[str, Any]:
    """Best-effort structured hints from resume text.

    Every value is optional. The UI shows them for review and the user decides
    what to keep, so a wrong guess never silently changes the profile.
    """
    cleaned = normalize_text(text or "")
    if not cleaned:
        return {}
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    emails = _EMAIL.findall(text or "")
    phones = [match.strip() for match in _PHONE.findall(text or "") if len(re.sub(r"\D", "", match)) >= 7]
    links: dict[str, str] = {}
    for pattern, key in ((_LINKEDIN, "linkedin"), (_GITHUB, "github")):
        found = pattern.search(text or "")
        if found:
            value = found.group(0)
            links[key] = value if value.lower().startswith("http") else f"https://{value}"
    for url in _URL.findall(text or "")[:6]:
        if "linkedin.com" in url.lower() or "github.com" in url.lower():
            continue
        links.setdefault("website", url.rstrip(".,;)"))
        break

    name = ""
    for line in lines[:6]:
        candidate = line.strip()
        if "@" in candidate or any(char.isdigit() for char in candidate):
            continue
        words = candidate.split()
        if 2 <= len(words) <= 5 and all(word.replace("-", "").isalpha() for word in words):
            name = candidate.title() if candidate.isupper() else candidate
            break

    headline = ""
    for line in lines[:10]:
        lowered = line.lower()
        if any(hint in lowered for hint in _HEADLINE_HINTS) and len(line) <= 120:
            headline = line
            break

    skills = extract_skills(cleaned, limit=30)
    years = 0.0
    for match in _YEARS.findall(cleaned):
        try:
            years = max(years, float(match))
        except ValueError:
            continue

    location = ""
    location_match = re.search(
        r"(?:location|based in|address)\s*[:\-]\s*([^\n|]{3,80})", text or "", re.I
    )
    if location_match:
        location = location_match.group(1).strip()
    if not location:
        for line in lines[:14]:
            if "@" in line or "http" in line.lower():
                continue
            resolved = resolve_location(line)
            if resolved.known and resolved.kind != "remote":
                location = resolved.display
                break

    summary = ""
    summary_match = re.search(
        r"(?:summary|profile|about me)\s*[:\-]?\s*\n?(.{40,600}?)(?:\n\s*\n|$)",
        text or "",
        re.I | re.S,
    )
    if summary_match:
        summary = re.sub(r"\s+", " ", summary_match.group(1)).strip()

    return {
        "full_name": name,
        "headline": headline,
        "email": emails[0] if emails else "",
        "phone": phones[0] if phones else "",
        "location": location,
        "links": links,
        "skills": skills,
        "years_experience": years,
        "summary": summary,
    }


class ResumeService:
    def __init__(self, db: Database | None = None, uploads_dir: Path | None = None) -> None:
        self.db = db or get_db()
        if uploads_dir is not None:
            self.uploads_dir = Path(uploads_dir)
        else:
            user_id = current_user_id()
            self.uploads_dir = (
                user_settings(user_id).uploads_dir if user_id else get_settings().uploads_dir
            )
        self.uploads_dir.mkdir(parents=True, exist_ok=True)

    # --------------------------------------------------------------------- read
    def list(self) -> list[Resume]:
        rows = self.db.query("SELECT * FROM resumes ORDER BY created_at DESC")
        return [self._row(row) for row in rows]

    def get(self, resume_id: str) -> Resume | None:
        row = self.db.query_one("SELECT * FROM resumes WHERE id = ?", (resume_id,))
        return self._row(row) if row else None

    def active(self) -> Resume | None:
        row = self.db.query_one(
            "SELECT * FROM resumes WHERE is_active = 1 ORDER BY updated_at DESC LIMIT 1"
        )
        return self._row(row) if row else None

    def file_path(self, resume_id: str) -> Path | None:
        row = self.db.query_one("SELECT stored_path FROM resumes WHERE id = ?", (resume_id,))
        if row is None:
            return None
        path = Path(row["stored_path"])
        return path if path.exists() else None

    # -------------------------------------------------------------------- write
    def upload(
        self,
        filename: str,
        data: bytes,
        *,
        mime_type: str = "",
        activate: bool = True,
    ) -> Resume:
        if not data:
            raise ResumeError("The uploaded file is empty.")
        if len(data) > MAX_RESUME_BYTES:
            raise ResumeError(
                f"Resume is larger than {MAX_RESUME_BYTES // (1024 * 1024)} MB."
            )
        cleaned = _clean_filename(filename)
        extension = _guess_extension(cleaned)
        if extension not in ALLOWED_EXTENSIONS:
            raise ResumeError("Upload a PDF, DOCX or plain-text resume.")
        for magic in _MAGIC[extension]:
            if not data.startswith(magic):
                raise ResumeError(
                    f"This file does not look like a valid {extension.lstrip('.').upper()}."
                )

        text, error = "", ""
        status = "ready"
        extracted: dict[str, Any] = {}
        try:
            text, error = _extract_text(extension, data)
            text = (text or "").strip()
            if not text:
                status = "error"
                error = (
                    "No selectable text was found. This looks like a scanned PDF. "
                    "Paste the text manually or upload a text-based file."
                )
            else:
                extracted = extract_fields(text)
        except ResumeError as exc:
            status = "error"
            error = str(exc)

        resume_id = _new_id()
        digest = hashlib.sha256(data).hexdigest()
        target = self.uploads_dir / f"{resume_id}{extension}"
        target.write_bytes(data)
        now = utc_now_iso()
        with self.db.transaction() as conn:
            if activate:
                conn.execute("UPDATE resumes SET is_active = 0")
            conn.execute(
                "INSERT INTO resumes(id, filename, stored_path, mime_type, size_bytes, sha256, "
                "status, text, extracted, error, is_active, created_at, updated_at) "
                "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                (
                    resume_id,
                    cleaned,
                    str(target),
                    mime_type or self._mime(extension),
                    len(data),
                    digest,
                    status,
                    text[:200_000],
                    json_dumps(extracted),
                    error,
                    int(activate),
                    now,
                    now,
                ),
            )
        stored = self.get(resume_id)
        if stored is None:  # pragma: no cover - defensive
            raise ResumeError("Could not store the uploaded resume.")
        return stored

    def activate(self, resume_id: str) -> Resume:
        if self.get(resume_id) is None:
            raise ResumeError("Resume not found.")
        with self.db.transaction() as conn:
            conn.execute("UPDATE resumes SET is_active = 0")
            conn.execute(
                "UPDATE resumes SET is_active = 1, updated_at = ? WHERE id = ?",
                (utc_now_iso(), resume_id),
            )
        stored = self.get(resume_id)
        assert stored is not None
        return stored

    def update_fields(self, resume_id: str, fields: dict[str, Any]) -> Resume:
        resume = self.get(resume_id)
        if resume is None:
            raise ResumeError("Resume not found.")
        merged = dict(resume.extracted)
        for key, value in fields.items():
            if key not in merged and key not in {
                "full_name",
                "headline",
                "email",
                "phone",
                "location",
                "links",
                "skills",
                "summary",
            }:
                continue
            if isinstance(value, str):
                merged[key] = value.strip()
            elif isinstance(value, (list, dict)):
                merged[key] = value
        self.db.execute(
            "UPDATE resumes SET extracted = ?, updated_at = ? WHERE id = ?",
            (json_dumps(merged), utc_now_iso(), resume_id),
        )
        updated = self.get(resume_id)
        assert updated is not None
        return updated

    def delete(self, resume_id: str) -> None:
        row = self.db.query_one("SELECT stored_path FROM resumes WHERE id = ?", (resume_id,))
        if row is None:
            raise ResumeError("Resume not found.")
        path = Path(row["stored_path"])
        self.db.execute("DELETE FROM resumes WHERE id = ?", (resume_id,))
        # A missing file must not block deleting the record.
        with suppress(OSError):
            path.unlink(missing_ok=True)

    # ----------------------------------------------------------------- helpers
    @staticmethod
    def _mime(extension: str) -> str:
        return {
            ".pdf": "application/pdf",
            ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            ".txt": "text/plain",
        }.get(extension, "application/octet-stream")

    @staticmethod
    def _row(row: Any) -> Resume:
        return Resume(
            id=row["id"],
            filename=row["filename"],
            mime_type=row["mime_type"] or "",
            size_bytes=int(row["size_bytes"] or 0),
            status=row["status"] or "uploaded",
            text=row["text"] or "",
            extracted=json_loads(row["extracted"], {}) or {},
            error=row["error"] or "",
            is_active=bool(row["is_active"]),
            created_at=row["created_at"] or "",
            updated_at=row["updated_at"] or "",
        )


def get_resumes() -> ResumeService:
    # Captures the active account, so it must not be cached across requests.
    return ResumeService()


__all__ = [
    "ALLOWED_EXTENSIONS",
    "MAX_RESUME_BYTES",
    "Resume",
    "ResumeError",
    "ResumeService",
    "extract_fields",
    "get_resumes",
]
